from __future__ import annotations

import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np

from src.utils.config import get_path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".jsonl", ".pdf", ".docx", ".pptx", ".svg"}
TURBOVEC_BACKENDS = {"turbovec", "turboquant", "turboquant-vector", "turboquant_vectors", "tq"}


class HashingEmbeddingModel:
    """Small dependency-light embedding fallback for constrained notebook runtimes."""

    def __init__(self, n_features: int = 1024) -> None:
        from sklearn.feature_extraction.text import HashingVectorizer

        self.vectorizer = HashingVectorizer(
            n_features=n_features,
            alternate_sign=False,
            norm="l2",
            stop_words="english",
        )

    def encode(self, texts: list[str], show_progress_bar: bool = False, normalize_embeddings: bool = True):
        del show_progress_bar, normalize_embeddings
        return self.vectorizer.transform(texts).astype("float32").toarray()


@dataclass
class RetrievedChunk:
    text: str
    source: str
    chunk_id: str
    score: float | None = None
    metadata: dict[str, Any] | None = None


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader
        except Exception:
            return ""
    try:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        print(f"Skipping unreadable PDF {path}: {exc}")
        return ""


def _read_docx(path: Path) -> str:
    try:
        import docx
    except Exception:
        return ""
    try:
        document = docx.Document(str(path))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception as exc:
        print(f"Skipping unreadable DOCX {path}: {exc}")
        return ""


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except Exception:
        return ""
    try:
        prs = Presentation(str(path))
        texts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
        return "\n".join(texts)
    except Exception as exc:
        print(f"Skipping unreadable PPTX {path}: {exc}")
        return ""


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".jsonl", ".svg"}:
        return _read_text_file(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pptx":
        return _read_pptx(path)
    return ""


def split_text(text: str, chunk_size: int = 900, chunk_overlap: int = 150) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(0, end - chunk_overlap)
    return chunks


class EducationalRAG:
    def __init__(self, cfg: dict[str, Any]) -> None:
        self.cfg = cfg
        self.rag_cfg = cfg.get("rag", {})
        self.index_dir = get_path(cfg, "rag_index_dir")
        self.docs_dir = get_path(cfg, "rag_docs_dir")
        self.retrieval_log = get_path(cfg, "retrieval_log")
        self.index_dir.mkdir(parents=True, exist_ok=True)
        self.retrieval_log.parent.mkdir(parents=True, exist_ok=True)
        self.backend = str(self.rag_cfg.get("backend", "chroma")).lower()
        self.collection_name = self.rag_cfg.get("collection_name", "ai_tutor_rag")
        self.embedding_model_name = self.rag_cfg.get("embedding_model", "sentence-transformers/all-MiniLM-L6-v2")
        self.embedding_backend = str(self.rag_cfg.get("embedding_backend", "auto")).lower()
        self.turbovec_bit_width = int(self.rag_cfg.get("turbovec_bits", self.rag_cfg.get("bit_width", 3)))
        self.turbovec_index_path = self.index_dir / f"{self.collection_name}.turbovec"
        self.turbovec_chunks_path = self.index_dir / f"{self.collection_name}_turbovec_chunks.json"
        self._embedding_model = None
        self._sentence_transformer_cls = None
        self._sentence_transformer_import_error: Exception | None = None
        self._client = None
        self._collection = None
        self._turbovec_index = None
        self._turbovec_chunks: list[dict[str, Any]] | None = None
        self._prime_embedding_import()

    def _uses_turbovec(self) -> bool:
        return self.backend in TURBOVEC_BACKENDS

    def _uses_hashing_embeddings(self) -> bool:
        return self.embedding_backend in {"hash", "hashing", "tfidf", "tfidf-hashing"} or self.embedding_model_name in {
            "hash",
            "hashing",
            "tfidf",
            "tfidf-hashing",
        }

    def _prime_embedding_import(self) -> None:
        if self._uses_hashing_embeddings():
            return
        try:
            from sentence_transformers import SentenceTransformer

            self._sentence_transformer_cls = SentenceTransformer
        except Exception as exc:
            self._sentence_transformer_import_error = exc

    @property
    def embedding_model(self):
        if self._embedding_model is None:
            if self._uses_hashing_embeddings():
                n_features = int(self.rag_cfg.get("hashing_features", 1024))
                print(f"Using HashingVectorizer RAG embeddings ({n_features} features)")
                self._embedding_model = HashingEmbeddingModel(n_features=n_features)
                return self._embedding_model
            try:
                print(f"Loading embedding model: {self.embedding_model_name}")
                if self._sentence_transformer_cls is None:
                    if self._sentence_transformer_import_error is not None:
                        raise self._sentence_transformer_import_error
                    from sentence_transformers import SentenceTransformer

                    self._sentence_transformer_cls = SentenceTransformer
                self._embedding_model = self._sentence_transformer_cls(self.embedding_model_name)
            except Exception as exc:
                if self.embedding_backend not in {"auto", "fallback"}:
                    raise
                n_features = int(self.rag_cfg.get("hashing_features", 1024))
                print(
                    "SentenceTransformer import/load failed; using HashingVectorizer fallback "
                    f"({n_features} features). Reason: {exc}"
                )
                self._embedding_model = HashingEmbeddingModel(n_features=n_features)
        return self._embedding_model

    @property
    def client(self):
        if self._client is None:
            import chromadb

            self._client = chromadb.PersistentClient(path=str(self.index_dir))
        return self._client

    def _get_or_create_collection(self):
        if self._collection is not None:
            return self._collection
        try:
            self._collection = self.client.get_collection(self.collection_name)
        except Exception:
            self._collection = self.client.create_collection(self.collection_name, metadata={"hnsw:space": "cosine"})
        return self._collection

    @staticmethod
    def _normalise_embeddings(embeddings: Any) -> np.ndarray:
        arr = np.asarray(embeddings, dtype=np.float32)
        if arr.ndim == 1:
            arr = arr.reshape(1, -1)
        norms = np.linalg.norm(arr, axis=1, keepdims=True)
        return arr / np.maximum(norms, 1e-12)

    def _load_turbovec_index(self) -> bool:
        if self._turbovec_index is not None and self._turbovec_chunks is not None:
            return True
        if not self.turbovec_index_path.exists() or not self.turbovec_chunks_path.exists():
            return False
        from turbovec import TurboQuantIndex

        self._turbovec_index = TurboQuantIndex.load(str(self.turbovec_index_path))
        with open(self.turbovec_chunks_path, "r", encoding="utf-8") as f:
            self._turbovec_chunks = json.load(f)
        return bool(self._turbovec_chunks)

    def has_existing_chunks(self) -> bool:
        if self._uses_turbovec():
            try:
                return self._load_turbovec_index()
            except Exception:
                return False
        try:
            collection = self._get_or_create_collection()
            return int(collection.count()) > 0
        except Exception:
            return False

    def load_documents(self) -> list[dict[str, Any]]:
        if not self.docs_dir.exists():
            print(f"RAG docs folder not found: {self.docs_dir}")
            return []
        max_file_mb = float(self.rag_cfg.get("max_file_mb", 30))
        max_files = int(self.rag_cfg.get("max_files", 0) or 0)
        max_chunks = int(self.rag_cfg.get("max_chunks", 0) or 0)
        max_chars_per_file = int(self.rag_cfg.get("max_chars_per_file", 0) or 0)
        chunk_size = int(self.rag_cfg.get("chunk_size", 900))
        chunk_overlap = int(self.rag_cfg.get("chunk_overlap", 150))
        docs: list[dict[str, Any]] = []
        files_seen = 0
        for path in sorted(self.docs_dir.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            if max_files and files_seen >= max_files:
                print(f"RAG file cap reached ({max_files} files).")
                break
            if path.stat().st_size > max_file_mb * 1024 * 1024:
                print(f"Skipping large RAG source {path} ({path.stat().st_size / (1024**2):.1f} MB)")
                continue
            text = read_document(path)
            if not text.strip():
                text = (
                    f"Document title: {path.stem.replace('_', ' ')}. "
                    f"Source file: {path.name}. Extension: {path.suffix.lower()}. "
                    "Body text was unavailable in this runtime."
                )
            if max_chars_per_file and len(text) > max_chars_per_file:
                text = text[:max_chars_per_file]
            files_seen += 1
            for i, chunk in enumerate(split_text(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)):
                docs.append(
                    {
                        "id": f"{path.stem}_{i}_{abs(hash(str(path))) % 10_000_000}",
                        "text": chunk,
                        "metadata": {"source": str(path), "chunk_index": i},
                    }
                )
                if max_chunks and len(docs) >= max_chunks:
                    print(f"RAG chunk cap reached ({max_chunks:,} chunks).")
                    print(f"Loaded {len(docs):,} RAG chunks from {self.docs_dir}")
                    return docs
        print(f"Loaded {len(docs):,} RAG chunks from {self.docs_dir}")
        return docs

    def build_index(self, force_rebuild: bool | None = None) -> None:
        force_rebuild = self.rag_cfg.get("rebuild_index", False) if force_rebuild is None else force_rebuild
        if self._uses_turbovec():
            self._build_turbovec_index(force_rebuild=bool(force_rebuild))
            return

        if force_rebuild:
            try:
                self.client.delete_collection(self.collection_name)
            except Exception:
                pass
            self._collection = self.client.create_collection(self.collection_name, metadata={"hnsw:space": "cosine"})
        collection = self._get_or_create_collection()
        if not force_rebuild:
            try:
                if collection.count() > 0:
                    print(f"Using existing RAG index at {self.index_dir} with {collection.count():,} chunks")
                    return
            except Exception:
                pass

        docs = self.load_documents()
        if not docs:
            print("No RAG documents found; retrieval will return empty context.")
            return
        batch_size = int(self.rag_cfg.get("embedding_batch_size", 64))
        for start in range(0, len(docs), batch_size):
            batch = docs[start : start + batch_size]
            texts = [d["text"] for d in batch]
            embeddings = self.embedding_model.encode(texts, show_progress_bar=False, normalize_embeddings=True).tolist()
            collection.add(
                ids=[d["id"] for d in batch],
                documents=texts,
                metadatas=[d["metadata"] for d in batch],
                embeddings=embeddings,
            )
            if start == 0 or (start + batch_size) % (batch_size * 10) == 0:
                print(f"Indexed {min(start + batch_size, len(docs)):,}/{len(docs):,} RAG chunks")
        meta_path = self.index_dir / "chunk_metadata.json"
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump([{"id": d["id"], **d["metadata"]} for d in docs], f, indent=2)
        print(f"Saved RAG chunk metadata: {meta_path}")

    def _build_turbovec_index(self, force_rebuild: bool = False) -> None:
        if not force_rebuild and self._load_turbovec_index():
            print(
                f"Using existing TurboVec RAG index at {self.turbovec_index_path} "
                f"with {len(self._turbovec_chunks or []):,} chunks"
            )
            return

        docs = self.load_documents()
        if not docs:
            print("No RAG documents found; retrieval will return empty context.")
            return

        from turbovec import TurboQuantIndex

        batch_size = int(self.rag_cfg.get("embedding_batch_size", 64))
        index = None
        stored_chunks: list[dict[str, Any]] = []
        for start in range(0, len(docs), batch_size):
            batch = docs[start : start + batch_size]
            texts = [d["text"] for d in batch]
            embeddings = self._normalise_embeddings(
                self.embedding_model.encode(texts, show_progress_bar=False, normalize_embeddings=True)
            )
            if index is None:
                index = TurboQuantIndex(dim=int(embeddings.shape[1]), bit_width=self.turbovec_bit_width)
            index.add(embeddings)
            stored_chunks.extend(
                {"id": d["id"], "text": d["text"], "metadata": d["metadata"]}
                for d in batch
            )
            if start == 0 or (start + batch_size) % (batch_size * 10) == 0:
                print(f"TurboVec indexed {min(start + batch_size, len(docs)):,}/{len(docs):,} RAG chunks")

        if index is None:
            print("No embeddings were created; retrieval will return empty context.")
            return
        try:
            index.prepare()
        except Exception:
            pass

        index.write(str(self.turbovec_index_path))
        with open(self.turbovec_chunks_path, "w", encoding="utf-8") as f:
            json.dump(stored_chunks, f, ensure_ascii=False, indent=2)
        meta_path = self.index_dir / "chunk_metadata.json"
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump([{"id": d["id"], **d["metadata"]} for d in docs], f, indent=2)
        self._turbovec_index = index
        self._turbovec_chunks = stored_chunks
        print(
            f"Saved TurboVec RAG index: {self.turbovec_index_path} "
            f"({len(stored_chunks):,} chunks, {self.turbovec_bit_width}-bit)"
        )

    def retrieve(
        self,
        student_question: str,
        weak_skill: str | None = None,
        lesson: str | None = None,
        top_k: int | None = None,
    ) -> tuple[list[RetrievedChunk], float]:
        self.build_index(force_rebuild=False)
        top_k = int(top_k or self.rag_cfg.get("top_k", 4))
        query = " ".join(x for x in [weak_skill or "", lesson or "", student_question] if x).strip()
        if not query:
            return [], 0.0
        if self._uses_turbovec():
            return self._retrieve_turbovec(query, top_k)

        collection = self._get_or_create_collection()
        if collection.count() == 0:
            return [], 0.0
        start = time.perf_counter()
        query_embedding = self.embedding_model.encode([query], show_progress_bar=False, normalize_embeddings=True).tolist()[0]
        result = collection.query(query_embeddings=[query_embedding], n_results=top_k, include=["documents", "metadatas", "distances"])
        latency_ms = (time.perf_counter() - start) * 1000
        chunks: list[RetrievedChunk] = []
        documents = result.get("documents", [[]])[0]
        metadatas = result.get("metadatas", [[]])[0]
        ids = result.get("ids", [[]])[0]
        distances = result.get("distances", [[]])[0]
        for doc, metadata, chunk_id, distance in zip(documents, metadatas, ids, distances):
            chunks.append(
                RetrievedChunk(
                    text=doc,
                    source=str((metadata or {}).get("source", "")),
                    chunk_id=str(chunk_id),
                    score=float(1.0 - distance) if distance is not None else None,
                    metadata=metadata or {},
                )
            )
        self._log_retrieval(query, chunks, latency_ms)
        return chunks, latency_ms

    def _retrieve_turbovec(self, query: str, top_k: int) -> tuple[list[RetrievedChunk], float]:
        if not self._load_turbovec_index():
            return [], 0.0
        chunks_meta = self._turbovec_chunks or []
        if not chunks_meta:
            return [], 0.0
        start = time.perf_counter()
        query_embedding = self._normalise_embeddings(
            self.embedding_model.encode([query], show_progress_bar=False, normalize_embeddings=True)
        )
        scores, indices = self._turbovec_index.search(query_embedding, k=min(top_k, len(chunks_meta)))
        latency_ms = (time.perf_counter() - start) * 1000
        chunks: list[RetrievedChunk] = []
        for score, idx in zip(scores[0], indices[0]):
            pos = int(idx)
            if pos < 0 or pos >= len(chunks_meta):
                continue
            item = chunks_meta[pos]
            metadata = dict(item.get("metadata") or {})
            metadata["rag_backend"] = "turbovec"
            chunks.append(
                RetrievedChunk(
                    text=str(item.get("text", "")),
                    source=str(metadata.get("source", "")),
                    chunk_id=str(item.get("id", pos)),
                    score=float(score),
                    metadata=metadata,
                )
            )
        self._log_retrieval(query, chunks, latency_ms)
        return chunks, latency_ms

    def _log_retrieval(self, query: str, chunks: list[RetrievedChunk], latency_ms: float) -> None:
        payload = {
            "query": query,
            "backend": self.backend,
            "latency_ms": latency_ms,
            "n_chunks": len(chunks),
            "chunks": [
                {"chunk_id": c.chunk_id, "source": c.source, "score": c.score, "preview": c.text[:160]}
                for c in chunks
            ],
        }
        with open(self.retrieval_log, "a", encoding="utf-8") as f:
            f.write(json.dumps(payload, ensure_ascii=False) + "\n")
